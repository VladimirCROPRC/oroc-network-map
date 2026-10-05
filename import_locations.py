import argparse
import json
import math
import shutil
import sqlite3
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


LAYER_ID = "oroc-locations"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def main():
    parser = argparse.ArgumentParser(description="Import OROC locations from Excel")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--fo-manifest", type=Path, required=True)
    parser.add_argument("--dist", type=Path, default=Path(__file__).resolve().parent / "dist")
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    sheet = workbook["Sheet2"]
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows)]
    required = {"COD_LOCATIE", "LATITUDINE", "LONGITUDINE"}
    if not required.issubset(headers):
        raise ValueError(f"Missing columns: {sorted(required - set(headers))}")

    index = {name: headers.index(name) for name in headers if name}
    temp_dir = Path(tempfile.mkdtemp(prefix="oroc-locations-", dir=args.dist.parent.resolve()))
    database = sqlite3.connect(temp_dir / "locations.sqlite")
    database.execute("PRAGMA journal_mode=OFF")
    database.execute("PRAGMA synchronous=OFF")
    database.execute("CREATE TABLE features (shard TEXT NOT NULL, payload TEXT NOT NULL)")
    counts = defaultdict(int)
    shard_bounds = {}
    total_bounds = [float("inf"), float("inf"), float("-inf"), float("-inf")]
    search_records = []
    count = skipped = 0

    try:
        for row in rows:
            code = clean(row[index["COD_LOCATIE"]])
            try:
                latitude = float(row[index["LATITUDINE"]])
                longitude = float(row[index["LONGITUDINE"]])
            except (TypeError, ValueError):
                skipped += 1
                continue
            if not code or not math.isfinite(latitude) or not math.isfinite(longitude) or latitude == 0 or longitude == 0:
                skipped += 1
                continue
            if not (43 <= latitude <= 49 and 20 <= longitude <= 30):
                skipped += 1
                continue

            point = [round(longitude, 6), round(latitude, 6)]
            key = f"{math.floor(longitude * 4)}_{math.floor(latitude * 4)}"
            bounds = shard_bounds.setdefault(key, [float("inf"), float("inf"), float("-inf"), float("-inf")])
            bounds[0], bounds[1] = min(bounds[0], longitude), min(bounds[1], latitude)
            bounds[2], bounds[3] = max(bounds[2], longitude), max(bounds[3], latitude)
            total_bounds[0], total_bounds[1] = min(total_bounds[0], longitude), min(total_bounds[1], latitude)
            total_bounds[2], total_bounds[3] = max(total_bounds[2], longitude), max(total_bounds[3], latitude)

            properties = {"Cod_Locatie": code}
            fields = {
                "DEN_LOCATIE": "Denumire",
                "JUDET": "Judet",
                "LOCALITATEA": "Localitate",
                "TIP_ART": "Tip_adresa",
                "STRADA": "Strada",
                "NR_IMOBIL": "Numar",
                "OBSERVATII": "Observatii",
            }
            for source, target in fields.items():
                value = clean(row[index[source]]) if source in index else ""
                if value:
                    properties[target] = value

            payload = json.dumps([properties, point], ensure_ascii=False, separators=(",", ":"))
            database.execute("INSERT INTO features (shard, payload) VALUES (?, ?)", (key, payload))
            counts[key] += 1
            count += 1
            search_records.append({"c": code, "n": properties.get("Denumire", code), "p": point})
            if count % 10000 == 0:
                database.commit()

        database.commit()
        database.execute("CREATE INDEX features_by_shard ON features (shard)")
        database.commit()

        layers_dir = args.dist / "layers"
        layers_dir.mkdir(parents=True, exist_ok=True)
        for old in layers_dir.glob(f"{LAYER_ID}-*.json"):
            old.unlink()
        shards = []
        for key in sorted(counts):
            filename = f"layers/{LAYER_ID}-{key}.json"
            output = args.dist / filename
            written = 0
            with output.open("w", encoding="utf-8") as target:
                target.write('{"p":[')
                first = True
                for (payload,) in database.execute("SELECT payload FROM features WHERE shard = ?", (key,)):
                    if not first:
                        target.write(",")
                    target.write(payload)
                    first = False
                    written += 1
                target.write("]}")
            if written != counts[key]:
                raise RuntimeError(f"Shard {key} contains {written} of {counts[key]} expected points")
            shards.append({"file": filename, "features": written, "bytes": output.stat().st_size, "bounds": shard_bounds[key]})
    finally:
        database.close()
        shutil.rmtree(temp_dir)

    if not count:
        raise ValueError("No valid locations found")

    search_dir = args.dist / "site-codes"
    search_dir.mkdir(parents=True, exist_ok=True)
    for old in search_dir.glob("*.json"):
        old.unlink()
    search_groups = defaultdict(list)
    for record in search_records:
        search_groups[record["c"][:2]].append(record)
    search_shards = []
    for key in sorted(search_groups):
        filename = f"site-codes/{key}.json"
        output = args.dist / filename
        output.write_text(json.dumps(search_groups[key], ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
        search_shards.append({"key": key, "file": filename, "records": len(search_groups[key]), "bytes": output.stat().st_size})
    search_path = args.dist / "site-codes.json"
    search_path.write_text(json.dumps({"records": len(search_records), "shards": search_shards}, separators=(",", ":")), encoding="utf-8")
    layer = {
        "id": LAYER_ID,
        "name": "Locații OROC",
        "group": "Locații",
        "features": count,
        "sourceFeatures": count + skipped,
        "skipped": skipped,
        "geometryTypes": {"Point": count},
        "bytes": sum(shard["bytes"] for shard in shards),
        "source": args.workbook.name,
        "bounds": total_bounds,
        "color": "#ffe500",
        "shards": shards,
    }
    source_manifest = json.loads(args.fo_manifest.read_text(encoding="utf-8"))
    fo_layer = next(item for item in source_manifest["layers"] if item["id"] == "fo-oroc")
    manifest = {
        "generatedFrom": args.workbook.name,
        "layers": [layer, fo_layer],
        "errors": [],
        "totals": {
            "layers": 2,
            "features": layer["features"] + fo_layer["features"],
            "skipped": layer["skipped"] + fo_layer["skipped"],
            "bytes": layer["bytes"] + fo_layer["bytes"],
        },
    }
    (args.dist / "layers.json").write_text(json.dumps(manifest, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"layer": layer, "searchBytes": sum(shard["bytes"] for shard in search_shards), "searchShards": len(search_shards)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
